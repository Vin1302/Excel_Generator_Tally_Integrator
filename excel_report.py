"""
excel_report.py
---------------
Build the output workbook with openpyxl:

  Sheet "Summary"   : A/B = Payment name + total  |  C/D = Received name + total
                      (sorted largest first). Each name is a hyperlink that jumps
                      to its group on the detail sheet.

  Sheet "Payments"  : one collapsible (+/-) group per payee. The group header row
                      shows the payee and total; the hidden child rows underneath
                      list every individual transaction with its date.

  Sheet "Receipts"  : same idea for money received.

The click-to-expand behaviour is Excel's native row outline grouping, so it works
in Excel, LibreOffice and Google Sheets with no macros.
"""

from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---- styling -------------------------------------------------------------- #
TITLE_FONT = Font(bold=True, size=14, color="FFFFFF")
HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
GROUP_FONT = Font(bold=True, size=11)
LINK_FONT = Font(color="0563C1", underline="single")
MONEY_FMT = "#,##0.00"
DATE_FMT = "dd-mmm-yyyy"

TITLE_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FILL = PatternFill("solid", fgColor="2E5496")
PAID_FILL = PatternFill("solid", fgColor="FCE4E4")
RECV_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_workbook(df: pd.DataFrame, out_path: str) -> str:
    """df must have columns: date, description, amount, direction, group."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    paid = df[df["direction"] == "paid"]
    recv = df[df["direction"] == "received"]

    # Build the detail sheets first so we know which row each group starts on
    # (needed for the Summary hyperlinks).
    paid_anchors = _build_detail_sheet(wb, "Payments", paid, PAID_FILL,
                                       "Money Paid Out")
    recv_anchors = _build_detail_sheet(wb, "Receipts", recv, RECV_FILL,
                                       "Money Received")

    _build_summary_sheet(summary, paid, recv, paid_anchors, recv_anchors)

    wb.save(out_path)
    return out_path


# --------------------------------------------------------------------------- #
def _group_totals(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty:
        return pd.DataFrame(columns=["group", "amount"])
    g = (frame.groupby("group", as_index=False)["amount"]
         .sum()
         .sort_values("amount", ascending=False)
         .reset_index(drop=True))
    return g


def _build_summary_sheet(ws, paid, recv, paid_anchors, recv_anchors):
    ws.merge_cells("A1:D1")
    c = ws["A1"]
    c.value = "Bank Statement Summary"
    c.font = TITLE_FONT
    c.fill = TITLE_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    headers = ["Payment Name", "Amount Paid", "Received From", "Amount Received"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER

    paid_tot = _group_totals(paid)
    recv_tot = _group_totals(recv)

    for r in range(len(paid_tot)):
        name = paid_tot.loc[r, "group"]
        amt = float(paid_tot.loc[r, "amount"])
        row = 3 + r
        cell = ws.cell(row=row, column=1, value=name)
        cell.border = BORDER
        if name in paid_anchors:
            cell.hyperlink = f"#Payments!A{paid_anchors[name]}"
            cell.font = LINK_FONT
        acell = ws.cell(row=row, column=2, value=amt)
        acell.number_format = MONEY_FMT
        acell.border = BORDER

    for r in range(len(recv_tot)):
        name = recv_tot.loc[r, "group"]
        amt = float(recv_tot.loc[r, "amount"])
        row = 3 + r
        cell = ws.cell(row=row, column=3, value=name)
        cell.border = BORDER
        if name in recv_anchors:
            cell.hyperlink = f"#Receipts!A{recv_anchors[name]}"
            cell.font = LINK_FONT
        acell = ws.cell(row=row, column=4, value=amt)
        acell.number_format = MONEY_FMT
        acell.border = BORDER

    # Totals line.
    total_row = 3 + max(len(paid_tot), len(recv_tot)) + 1
    tp = ws.cell(row=total_row, column=1, value="TOTAL PAID")
    tp.font = GROUP_FONT
    tpa = ws.cell(row=total_row, column=2, value=float(paid_tot["amount"].sum()))
    tpa.number_format = MONEY_FMT
    tpa.font = GROUP_FONT
    tr = ws.cell(row=total_row, column=3, value="TOTAL RECEIVED")
    tr.font = GROUP_FONT
    tra = ws.cell(row=total_row, column=4, value=float(recv_tot["amount"].sum()))
    tra.number_format = MONEY_FMT
    tra.font = GROUP_FONT

    _autosize(ws, {1: 34, 2: 16, 3: 34, 4: 16})
    ws.freeze_panes = "A3"


def _build_detail_sheet(wb, title, frame, fill, banner):
    """Returns {group_name: header_row_number} for hyperlink targets."""
    ws = wb.create_sheet(title)
    ws.merge_cells("A1:C1")
    c = ws["A1"]
    c.value = banner
    c.font = TITLE_FONT
    c.fill = TITLE_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    for i, h in enumerate(["Date", "Details", "Amount"], start=1):
        cell = ws.cell(row=2, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER

    # Summary (group header) row sits ABOVE its detail rows.
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.applyStyles = True

    anchors = {}
    row = 3

    totals = _group_totals(frame)
    for _, g in totals.iterrows():
        gname = g["group"]
        gtotal = float(g["amount"])

        # --- group header row (always visible, level 0) ---
        hcell = ws.cell(row=row, column=1, value=gname)
        hcell.font = GROUP_FONT
        hcell.fill = fill
        ws.cell(row=row, column=2, value="").fill = fill
        tcell = ws.cell(row=row, column=3, value=gtotal)
        tcell.number_format = MONEY_FMT
        tcell.font = GROUP_FONT
        tcell.fill = fill
        anchors[gname] = row
        row += 1

        # --- detail rows (grouped + hidden -> collapsed by default) ---
        rows = frame[frame["group"] == gname].sort_values("date")
        for _, t in rows.iterrows():
            dcell = ws.cell(row=row, column=1, value=pd.to_datetime(t["date"]))
            dcell.number_format = DATE_FMT
            ws.cell(row=row, column=2, value=str(t["description"]))
            acell = ws.cell(row=row, column=3, value=float(t["amount"]))
            acell.number_format = MONEY_FMT
            ws.row_dimensions[row].outline_level = 1
            ws.row_dimensions[row].hidden = True   # start collapsed
            row += 1

    _autosize(ws, {1: 15, 2: 52, 3: 16})
    ws.freeze_panes = "A3"
    return anchors


def _autosize(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
